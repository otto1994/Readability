# -*- coding: utf-8 -*-

i = input("Please feed me row number : ")
name = input("Please tell me filename : ")
id = str(name)

from docx import Document
document = Document(id+".docx")

'''
這是可以用的版本，但是我仍然必需弄清楚逃脫符號的用法，
為何\n\n會在句子判斷上出現大量超估的情況，我認為應該是在斷句的判定上把換行認為是斷句，
而\n將換行逃脫，然而不無視換行，這個意思是換行後的尾頭兩字不會被黏在一起。
'''
docText = '\n'.join([
    paragraph.text.encode('utf-8') for paragraph in document.paragraphs
])

'''
問題在於對於item的判定:如果找不到item則返回""
'''

n = docText.count('@')+1

'''
對n的判定決定編碼能否運行，規則上應該是@的數目+1
'''

import re
sp = re.split(r'@*',docText)

##################
for j in range(0,n):
	if sp[j] == "item1":
		item1 = sp[j+1]
	if "item1" not in sp:
		item1 = ""
##################
for j in range(0,n):
    if sp[j] == "item1a":
        item1a = sp[j+1]
    if "item1a" not in sp:
        item1a = ""
##################
for j in range(0,n):
    if sp[j] == "item1b":
        item1b = sp[j+1]
    if "item1b" not in sp:
        item1b = ""
##################
for j in range(0,n):
    if sp[j] == "item2":
        item2 = sp[j+1]
    if "item2" not in sp:
        item2 = ""
##################
for j in range(0,n):
    if sp[j] == "item3":
        item3 = sp[j+1]
    if "item3" not in sp:
        item3 = ""
##################
for j in range(0,n):
    if sp[j] == "item4":
        item4 = sp[j+1]
    if "item4" not in sp:
        item4 = ""
##################
for j in range(0,n):
    if sp[j] == "item4a":
        item4a = sp[j+1]
    if "item4a" not in sp:
        item4a = ""
##################
for j in range(0,n):
    if sp[j] == "item5":
        item5 = sp[j+1]
    if "item5" not in sp:
        item5 = ""
##################
for j in range(0,n):
    if sp[j] == "item6":
        item6 = sp[j+1]
    if "item6" not in sp:
        item6 = ""
##################
for j in range(0,n):
    if sp[j] == "item7":
        item7 = sp[j+1]
    if "item7" not in sp:
        item7 = ""
###################
for j in range(0,n):
    if sp[j] == "ResultsOfOperations":
        mda = sp[j+1]
    if "ResultsOfOperations" not in sp:
        mda = ""
###################
for j in range(0,n):
    if sp[j] == "item7a":
        item7a = sp[j+1]
    if "item7a" not in sp:
        item7a = ""
###################
for j in range(0,n):
    if sp[j] == "item8":
        item8 = sp[j+1]
    if "item8" not in sp:
        item8 = ""
###################
for j in range(0,n):
    if sp[j] == "item9":
        item9 = sp[j+1]
    if "item9" not in sp:
        item9 = ""
###################
for j in range(0,n):
    if sp[j] == "item9a":
        item9a = sp[j+1]
    if "item9a" not in sp:
        item9a = ""
###################
for j in range(0,n):
    if sp[j] == "item9b":
        item9b = sp[j+1]
    if "item9b" not in sp:
        item9b = ""
###################
for j in range(0,n):
    if sp[j] == "item10":
        item10to14 = sp[j+1]
    if "item10" not in sp:
        item10to14 = ""
#############################settings


import requests
import os.path
from xlwt import Workbook
from xlrd import open_workbook
from xlutils.copy import copy

if os.path.exists("record.xls")!= True:
    wb = Workbook()
    sheet1 = wb.add_sheet('Sheet1')

else:
    rb = open_workbook('record.xls')
    wb = copy(rb)
    sheet1 = wb.get_sheet(0)

###################
#below are codes for the items respectively.
#########################code for item1 
payload = {	
"type":"text",
"text":item1
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,0,id)
sheet1.write(i,1,f)
sheet1.write(i,2,g)
sheet1.write(i,3,c)
sheet1.write(i,4,s)
sheet1.write(i,5,a)
sheet1.write(i,6,average)
sheet1.write(i,7,words)
sheet1.write(i,8,sentences)

######################code for item1a
payload = {	
"type":"text",
"text":item1a
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,9,f)
sheet1.write(i,10,g)
sheet1.write(i,11,c)
sheet1.write(i,12,s)
sheet1.write(i,13,a)
sheet1.write(i,14,average)
sheet1.write(i,15,words)
sheet1.write(i,16,sentences)


############################code for item1b
payload = {	
"type":"text",
"text":item1b
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')



sheet1.write(i,17,f)
sheet1.write(i,18,g)
sheet1.write(i,19,c)
sheet1.write(i,20,s)
sheet1.write(i,21,a)
sheet1.write(i,22,average)
sheet1.write(i,23,words)
sheet1.write(i,24,sentences)

###########################code for item2
payload = {	
"type":"text",
"text":item2
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,25,f)
sheet1.write(i,26,g)
sheet1.write(i,27,c)
sheet1.write(i,28,s)
sheet1.write(i,29,a)
sheet1.write(i,30,average)
sheet1.write(i,31,words)
sheet1.write(i,32,sentences)

##############################code for item3
payload = {	
"type":"text",
"text":item3
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')

sheet1.write(i,33,f)
sheet1.write(i,34,g)
sheet1.write(i,35,c)
sheet1.write(i,36,s)
sheet1.write(i,37,a)
sheet1.write(i,38,average)
sheet1.write(i,39,words)
sheet1.write(i,40,sentences)

###########################code for item4
payload = {	
"type":"text",
"text":item4
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')

sheet1.write(i,41,f)
sheet1.write(i,42,g)
sheet1.write(i,43,c)
sheet1.write(i,44,s)
sheet1.write(i,45,a)
sheet1.write(i,46,average)
sheet1.write(i,47,words)
sheet1.write(i,48,sentences)

###########################code for item4a
payload = {	
"type":"text",
"text":item4a
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,49,f)
sheet1.write(i,50,g)
sheet1.write(i,51,c)
sheet1.write(i,52,s)
sheet1.write(i,53,a)
sheet1.write(i,54,average)
sheet1.write(i,55,words)
sheet1.write(i,56,sentences)

############################code for item5
payload = {	
"type":"text",
"text":item5
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')

sheet1.write(i,57,f)
sheet1.write(i,58,g)
sheet1.write(i,59,c)
sheet1.write(i,60,s)
sheet1.write(i,61,a)
sheet1.write(i,62,average)
sheet1.write(i,63,words)
sheet1.write(i,64,sentences)

##########################code for item6
payload = {	
"type":"text",
"text":item6
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,65,f)
sheet1.write(i,66,g)
sheet1.write(i,67,c)
sheet1.write(i,68,s)
sheet1.write(i,69,a)
sheet1.write(i,70,average)
sheet1.write(i,71,words)
sheet1.write(i,72,sentences)

#########################code for item7
payload = {	
"type":"text",
"text":item7
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,73,f)
sheet1.write(i,74,g)
sheet1.write(i,75,c)
sheet1.write(i,76,s)
sheet1.write(i,77,a)
sheet1.write(i,78,average)
sheet1.write(i,79,words)
sheet1.write(i,80,sentences)

##############################code for mda
payload = {	
"type":"text",
"text":mda
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,81,f)
sheet1.write(i,82,g)
sheet1.write(i,83,c)
sheet1.write(i,84,s)
sheet1.write(i,85,a)
sheet1.write(i,86,average)
sheet1.write(i,87,words)
sheet1.write(i,88,sentences)

#############################code for item7a

payload = {	
"type":"text",
"text":item7a
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,89,f)
sheet1.write(i,90,g)
sheet1.write(i,91,c)
sheet1.write(i,92,s)
sheet1.write(i,93,a)
sheet1.write(i,94,average)
sheet1.write(i,95,words)
sheet1.write(i,96,sentences)

################################code for item7+item7a( md&a df2)
payload = {	
"type":"text",
"text":item7+mda+item7a
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,97,f)
sheet1.write(i,98,g)
sheet1.write(i,99,c)
sheet1.write(i,100,s)
sheet1.write(i,101,a)
sheet1.write(i,102,average)
sheet1.write(i,103,words)
sheet1.write(i,104,sentences)

####################################code for item8 (note)
payload = {	
"type":"text",
"text":item8
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,105,f)
sheet1.write(i,106,g)
sheet1.write(i,107,c)
sheet1.write(i,108,s)
sheet1.write(i,109,a)
sheet1.write(i,110,average)
sheet1.write(i,111,words)
sheet1.write(i,112,sentences)

###################################code for item9
payload = {	
"type":"text",
"text":item9
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,113,f)
sheet1.write(i,114,g)
sheet1.write(i,115,c)
sheet1.write(i,116,s)
sheet1.write(i,117,a)
sheet1.write(i,118,average)
sheet1.write(i,119,words)
sheet1.write(i,120,sentences)

##############################code for item9a
payload = {	
"type":"text",
"text":item9a
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,121,f)
sheet1.write(i,122,g)
sheet1.write(i,123,c)
sheet1.write(i,124,s)
sheet1.write(i,125,a)
sheet1.write(i,126,average)
sheet1.write(i,127,words)
sheet1.write(i,128,sentences)

################################code for item9b
payload = {	
"type":"text",
"text":item9b
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,129,f)
sheet1.write(i,130,g)
sheet1.write(i,131,c)
sheet1.write(i,132,s)
sheet1.write(i,133,a)
sheet1.write(i,134,average)
sheet1.write(i,135,words)
sheet1.write(i,136,sentences)

###############################code for item10to14
payload = {	
"type":"text",
"text":item10to14
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')


sheet1.write(i,137,f)
sheet1.write(i,138,g)
sheet1.write(i,139,c)
sheet1.write(i,140,s)
sheet1.write(i,141,a)
sheet1.write(i,142,average)
sheet1.write(i,143,words)
sheet1.write(i,144,sentences)

#############################code for overall counting
payload = {	
"type":"text",
"text":item1+item1a+item1b+item2+item3+item4+item4a+item5+item6+item7+mda+item7a+item8+item9+item9a+item9b+item10to14
}
res = requests.post("https://readability-score.com/ajax.php",data=payload)
text = str(res.text)
split = re.split(r'"*',text)

fkg = split[4].replace(':','')
f = fkg.replace(',','')
gf = split[6].replace(':','')
g = gf.replace(',','')
cl = split[8].replace(':','')
c = cl.replace(',','')
si = split[10].replace(':','')
s = si.replace(',','')
ar = split[12].replace(':','')
a = ar.replace(',','')
ag = split[15].replace(':','')
average = ag.replace(',','')
words = split[27].replace(',','')
sentences = split[19].replace(',','')

sheet1.write(i,145,f)
sheet1.write(i,146,g)
sheet1.write(i,147,c)
sheet1.write(i,148,s)
sheet1.write(i,149,a)
sheet1.write(i,150,average)
sheet1.write(i,151,words)
sheet1.write(i,152,sentences)



#####################finally, save the work
wb.save('record.xls')

###############hmmmm....
print 'comeon'